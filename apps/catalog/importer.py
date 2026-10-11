"""Import a course catalog (categories, courses, chapters, lessons, bundles) from an Excel file.

The parser is deliberately tolerant: headers are matched by Persian names (ignoring ZWNJ,
parentheses and digit style), sheets by their role (دسته‌ها / دوره‌ها / پکیج‌ها / درس‌های X /
تنظیمات قیمت). Everything is validated first; nothing is written unless the whole file is valid.
Re-importing updates existing rows (matched by course code / lesson code) instead of duplicating.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field

from django.core.files.base import ContentFile
from django.db import transaction

from .models import Category, Chapter, Course, Lesson, _unique_slug

MAX_COVER_BYTES = 5 * 1024 * 1024
MAX_ZIP_FILES = 500
MAX_ZIP_TOTAL_BYTES = 150 * 1024 * 1024
LIFETIME_DAYS = 36500
COVER_EXTENSIONS = ("png", "jpg", "jpeg", "webp")

_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


# ── cell helpers ─────────────────────────────────────────────────────────────
def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _key(value) -> str:
    text = _s(value).translate(_DIGITS)
    text = text.replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def _header_key(value) -> str:
    return re.sub(r"\s*\(.*?\)", "", _key(value)).strip()


def _num(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _s(value).translate(_DIGITS)
    text = re.sub(r"[,٬،\s]|تومان", "", text).replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.strip().startswith("=")


def _split_codes(value) -> list[str]:
    text = _s(value).translate(_DIGITS)
    if not text or text.strip() in {"-", "—"}:
        return []
    parts = re.split(r"[،,;/\n]+|\s+(?:و|یا|or|and)\s+|\s+", text)
    return [p.strip().upper() for p in parts if p.strip() and p.strip() not in {"و", "یا"}]


def _round(value: float, step: int) -> int:
    return int((value + step / 2) // step * step)


# ── column aliases (normalised with _header_key) ─────────────────────────────
def _aliases(**fields: list[str]) -> dict[str, set[str]]:
    return {name: {_header_key(x) for x in names} for name, names in fields.items()}


CATEGORY_COLS = _aliases(
    code=["کد دسته", "کد"],
    title=["نام دسته", "عنوان دسته", "عنوان", "دسته"],
    description=["توضیح", "توضیحات"],
    order=["ترتیب"],
)
COURSE_COLS = _aliases(
    code=["کد", "کد دوره"],
    title=["عنوان دوره", "عنوان", "نام دوره"],
    category=["دسته", "دسته بندی", "کد دسته"],
    kind=["نوع"],
    level=["سطح"],
    prereq=["پیش نیاز", "پیش نیازها"],
    hours=["مدت ویدیو"],
    tier=["رده قیمت"],
    price=["قیمت"],
    access=["دسترسی"],
    free=["درس رایگان", "رایگان"],
    status=["وضعیت"],
    summary=["خلاصه"],
    description=["توضیحات", "توضیح", "شرح"],
    cover=["تصویر", "تصویر جلد", "کاور"],
    slug=["نامک", "slug"],
    order=["ترتیب"],
)
BUNDLE_COLS = _aliases(
    code=["کد"],
    title=["نام پکیج", "عنوان پکیج", "عنوان"],
    courses=["دوره ها", "کد دوره ها"],
    discount=["تخفیف"],
    price=["قیمت پکیج", "قیمت"],
    access=["دسترسی"],
    status=["وضعیت"],
    summary=["خلاصه"],
    description=["توضیحات", "توضیح", "شرح"],
    cover=["تصویر", "تصویر جلد", "کاور"],
    slug=["نامک", "slug"],
    order=["ترتیب"],
)
LESSON_COLS = _aliases(
    course=["دوره", "کد دوره"],
    ch_no=["شماره فصل"],
    ch_title=["فصل", "عنوان فصل", "نام فصل"],
    code=["کد درس", "شماره درس", "شماره", "کد"],
    title=["عنوان درس", "نام درس", "عنوان"],
    kind=["نوع"],
    minutes=["مدت"],
    free=["رایگان"],
    body=["توضیح", "توضیحات", "متن"],
    video=["نام فایل ویدیو", "فایل ویدیو", "ویدیو"],
    files=["فایل های پیوست", "پیوست ها", "فایل ها"],
    note=["یادداشت"],
    status=["وضعیت"],
)

PUBLISHED_STATUSES = {
    _key(x)
    for x in [
        "آماده",
        "محتوا آماده",
        "آپلود شد",
        "منتشر",
        "منتشر شده",
        "published",
        "ready",
        "live",
    ]
}
FREE_WORDS = {_key(x) for x in ["رایگان", "free"]}
PERMANENT_WORDS = {_key(x) for x in ["نامحدود", "دائمی", "مادام العمر", "مادام‌العمر", "-", "—"]}
YES = {_key(x) for x in ["بله", "yes", "true", "1", "x", "✓", "رایگان"]}
NO = {_key(x) for x in ["خیر", "نه", "no", "false", "0"]}

LEVELS = {
    "مقدماتی": Course.Level.BEGINNER,
    "beginner": Course.Level.BEGINNER,
    "متوسط": Course.Level.INTERMEDIATE,
    "intermediate": Course.Level.INTERMEDIATE,
    "متوسط پیشرفته": Course.Level.INTERMEDIATE,
    "متوسط تا پیشرفته": Course.Level.INTERMEDIATE,
    "پیشرفته": Course.Level.ADVANCED,
    "advanced": Course.Level.ADVANCED,
}
KINDS = {
    "ویدیو": Lesson.Kind.VIDEO,
    "video": Lesson.Kind.VIDEO,
    "ارائه html": Lesson.Kind.DOCUMENT,
    "ارائه": Lesson.Kind.DOCUMENT,
    "اسلاید": Lesson.Kind.DOCUMENT,
    "pdf": Lesson.Kind.DOCUMENT,
    "pdf / اسلاید": Lesson.Kind.DOCUMENT,
    "document": Lesson.Kind.DOCUMENT,
    "نوت بوک": Lesson.Kind.NOTEBOOK,
    "notebook": Lesson.Kind.NOTEBOOK,
    "داده مکانی": Lesson.Kind.GEODATA,
    "geodata": Lesson.Kind.GEODATA,
    "کد": Lesson.Kind.CODE,
    "code": Lesson.Kind.CODE,
    "متن": Lesson.Kind.TEXT,
    "text": Lesson.Kind.TEXT,
}


# ── report ───────────────────────────────────────────────────────────────────
@dataclass
class ImportReport:
    dry_run: bool = False
    created: Counter = field(default_factory=Counter)
    updated: Counter = field(default_factory=Counter)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def lines(self) -> list[str]:
        out = []
        for label in ("دسته", "دوره", "فصل", "جلسه", "پکیج", "تصویر"):
            new, old = self.created[label], self.updated[label]
            if new or old:
                out.append(f"{label}: {new} جدید، {old} به‌روزرسانی")
        return out


class ImportErrorList(Exception):
    pass


# ── covers source ────────────────────────────────────────────────────────────
class CoverSource:
    """Images by filename, from a directory or a zip (matched by basename, case-insensitive)."""

    def __init__(self, files: dict[str, bytes] | None = None):
        self._files = {k.lower(): v for k, v in (files or {}).items()}

    @classmethod
    def from_zip(cls, fileobj) -> CoverSource:
        files: dict[str, bytes] = {}
        with zipfile.ZipFile(fileobj) as archive:
            infos = [i for i in archive.infolist() if not i.is_dir()]
            if len(infos) > MAX_ZIP_FILES or sum(i.file_size for i in infos) > MAX_ZIP_TOTAL_BYTES:
                raise ValueError("فایل zip بیش از حد بزرگ است.")
            for info in infos:
                name = info.filename.replace("\\", "/").rsplit("/", 1)[-1]
                ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
                if ext in COVER_EXTENSIONS and info.file_size <= MAX_COVER_BYTES:
                    files[name] = archive.read(info)
        return cls(files)

    @classmethod
    def from_dir(cls, path) -> CoverSource:
        from pathlib import Path

        files = {}
        for item in Path(path).iterdir():
            if item.suffix.lstrip(".").lower() in COVER_EXTENSIONS and item.is_file():
                if item.stat().st_size <= MAX_COVER_BYTES:
                    files[item.name] = item.read_bytes()
        return cls(files)

    def find(self, name: str = "", code: str = "") -> tuple[str, bytes] | None:
        candidates = []
        if name:
            candidates.append(name.rsplit("/", 1)[-1].lower())
        if code:
            candidates += [f"{code.lower()}.{ext}" for ext in COVER_EXTENSIONS]
        for candidate in candidates:
            if candidate in self._files:
                return candidate, self._files[candidate]
        return None


# ── parsed rows ──────────────────────────────────────────────────────────────
@dataclass
class CategoryRow:
    code: str
    title: str
    description: str = ""
    order: int = 0


@dataclass
class CourseRow:
    where: str
    code: str
    title: str
    category: str
    level: str
    prereqs: list[str]
    price: int
    access_days: int | None
    free_rule: tuple
    published: bool
    summary: str
    description: str
    cover: str
    slug: str
    order: int


@dataclass
class LessonRow:
    where: str
    course: str
    ch_no: int
    ch_title: str
    code: str
    title: str
    kind: str
    minutes: int | None
    free: bool | None
    body: str
    note: str
    published: bool


@dataclass
class BundleRow:
    where: str
    code: str
    title: str
    courses: list[str]
    price: int
    access_days: int | None
    published_status: bool
    summary: str
    description: str
    cover: str
    slug: str
    order: int
    discount: float = 0


def _cell(row, cols, name):
    index = cols.get(name)
    return row[index] if index is not None and index < len(row) else None


def _read_table(ws, aliases):
    rows = list(ws.iter_rows(values_only=True))
    for header_index, row in enumerate(rows[:6]):
        cols: dict[str, int] = {}
        for col_index, value in enumerate(row):
            header = _header_key(value)
            if not header:
                continue
            for name, names in aliases.items():
                if header in names and name not in cols:
                    cols[name] = col_index
        if len(cols) >= 2:
            body = [(header_index + 2 + i, r) for i, r in enumerate(rows[header_index + 1 :])]
            return cols, body
    return None, []


def _sheet_role(title: str) -> str | None:
    key = _key(title)
    for prefix, role in (
        ("درس", "lessons"),
        ("تنظیمات", "settings"),
        ("پکیج", "bundles"),
        ("دوره", "courses"),
        ("دسته", "categories"),
    ):
        if key.startswith(prefix):
            return role
    return None


def _parse_free_rule(text: str):
    key = _key(text)
    if key in {"", "-", "—", "هیچ", "ندارد", "هیچکدام"}:
        return ("none",)
    if key in {"همه", "all"}:
        return ("all",)
    if match := re.search(r"فصل\s*(\d+)", key):
        return ("chapter", int(match.group(1)))
    if match := re.search(r"(\d+)\s*درس", key):
        return ("first", int(match.group(1)))
    return None


def _parse_access(value, *, free: bool, default_days: int | None):
    """Days of access, or None for permanent."""
    if _key(value) in PERMANENT_WORDS and _s(value):
        return None
    number = _num(value)
    if number is not None and number > 0:
        return int(number)
    if free:
        return None
    return default_days if default_days is not None else 365


# ── parsing ──────────────────────────────────────────────────────────────────
class _Parsed:
    def __init__(self):
        self.categories: dict[str, CategoryRow] = {}
        self.courses: list[CourseRow] = []
        self.lessons: list[LessonRow] = []
        self.bundles: list[BundleRow] = []
        self.base_price: float | None = None
        self.tiers: dict[str, float] = {}
        self.default_access: int | None = None


def _parse_settings(ws, parsed: _Parsed):
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        key = _key(row[0])
        value = _num(row[1]) if len(row) > 1 else None
        if value is None:
            continue
        if key.startswith("قیمت پایه"):
            parsed.base_price = value
        elif key.startswith("ضریب رده"):
            parsed.tiers[key[len("ضریب رده") :].strip()] = value
        elif key.startswith("دسترسی"):
            parsed.default_access = int(value)


def _parse_categories(ws, parsed: _Parsed, report: ImportReport):
    cols, rows = _read_table(ws, CATEGORY_COLS)
    if cols is None or "title" not in cols:
        report.errors.append(f"برگه «{ws.title}»: ستون «نام دسته» پیدا نشد.")
        return
    for _number, row in rows:
        title = _s(_cell(row, cols, "title"))
        if not title:
            continue
        code = _s(_cell(row, cols, "code")).upper()
        order = _num(_cell(row, cols, "order"))
        item = CategoryRow(
            code=code,
            title=title,
            description=_s(_cell(row, cols, "description")),
            order=int(order) if order is not None else len(parsed.categories) + 1,
        )
        parsed.categories[_key(code or title)] = item
        parsed.categories[_key(title)] = item


def _parse_courses(ws, parsed: _Parsed, report: ImportReport):
    cols, rows = _read_table(ws, COURSE_COLS)
    if cols is None or "title" not in cols:
        report.errors.append(f"برگه «{ws.title}»: ستون «عنوان دوره» پیدا نشد.")
        return
    for number, row in rows:
        where = f"برگه «{ws.title}» ردیف {number}"
        title = _s(_cell(row, cols, "title"))
        category = _s(_cell(row, cols, "category"))
        if not title or not category:
            continue  # blank or totals row
        code = _s(_cell(row, cols, "code")).upper()
        level_raw = re.sub(r"\s*[-–—]\s*", " ", _key(_cell(row, cols, "level")))
        level = LEVELS.get(level_raw)
        if level is None:
            report.errors.append(f"{where}: سطح «{_s(_cell(row, cols, 'level'))}» نامعتبر است.")
            continue
        kind = _key(_cell(row, cols, "kind"))
        tier = _key(_cell(row, cols, "tier"))
        raw_price = _cell(row, cols, "price")
        price = _num(raw_price)
        if price is None:
            if kind in FREE_WORDS:
                price = 0
            elif tier in parsed.tiers and parsed.base_price:
                price = parsed.tiers[tier] * parsed.base_price
            elif tier in {"", "-", "—"} and kind not in {"پولی", "paid"}:
                price = 0
            else:
                report.errors.append(f"{where}: قیمت مشخص نیست (عدد یا «رده قیمت» لازم است).")
                continue
        elif kind in FREE_WORDS and price > 0:
            report.warnings.append(f"{where}: نوع «رایگان» است ولی قیمت {int(price)} وارد شده.")
        price = int(round(price))
        free_text = _s(_cell(row, cols, "free"))
        rule = _parse_free_rule(free_text)
        if rule is None:
            report.warnings.append(
                f"{where}: «درس رایگان» ({free_text}) فهمیده نشد؛ نادیده گرفته شد."
            )
            rule = ("none",)
        order = _num(_cell(row, cols, "order"))
        parsed.courses.append(
            CourseRow(
                where=where,
                code=code,
                title=title,
                category=category,
                level=level,
                prereqs=_split_codes(_cell(row, cols, "prereq")),
                price=price,
                access_days=_parse_access(
                    _cell(row, cols, "access"), free=price == 0, default_days=parsed.default_access
                ),
                free_rule=rule,
                published=_key(_cell(row, cols, "status")) in PUBLISHED_STATUSES,
                summary=_s(_cell(row, cols, "summary"))[:300],
                description=_s(_cell(row, cols, "description")),
                cover=_s(_cell(row, cols, "cover")),
                slug=_s(_cell(row, cols, "slug")),
                order=int(order) if order is not None else len(parsed.courses) + 1,
            )
        )


def _parse_lessons(ws, parsed: _Parsed, report: ImportReport):
    cols, rows = _read_table(ws, LESSON_COLS)
    if cols is None or "title" not in cols:
        report.errors.append(f"برگه «{ws.title}»: ستون «عنوان درس» پیدا نشد.")
        return
    sheet_course = ""
    if match := re.search(r"\b([a-z]+\d+)\b", _key(ws.title)):
        sheet_course = match.group(1).upper()
    chapter_numbers: dict[tuple[str, str], int] = {}
    for number, row in rows:
        where = f"برگه «{ws.title}» ردیف {number}"
        title = _s(_cell(row, cols, "title"))
        ch_title = _s(_cell(row, cols, "ch_title"))
        raw_ch_no = _num(_cell(row, cols, "ch_no"))
        if not title or (raw_ch_no is None and not ch_title):
            continue  # blank or totals row
        course = _s(_cell(row, cols, "course")).upper() or sheet_course
        if not course:
            report.errors.append(f"{where}: کد دوره مشخص نیست (ستون «دوره» یا نام برگه).")
            continue
        if raw_ch_no is not None:
            ch_no = int(raw_ch_no)
        elif match := re.fullmatch(r"فصل\s*(\d+)", _key(ch_title)):
            ch_no = int(match.group(1))
        else:
            ch_no = chapter_numbers.setdefault(
                (course, _key(ch_title)),
                1 + sum(1 for k in chapter_numbers if k[0] == course),
            )
        kind_text = _key(_cell(row, cols, "kind"))
        video = _s(_cell(row, cols, "video"))
        if kind_text:
            kind = KINDS.get(kind_text)
            if kind is None:
                report.errors.append(f"{where}: نوع جلسه «{kind_text}» نامعتبر است.")
                continue
        else:
            kind = Lesson.Kind.VIDEO if video or "video" not in cols else Lesson.Kind.DOCUMENT
        free_text = _key(_cell(row, cols, "free"))
        free = True if free_text in YES else False if free_text in NO else None
        minutes = _num(_cell(row, cols, "minutes"))
        status = _key(_cell(row, cols, "status"))
        note_parts = []
        if video:
            note_parts.append(f"ویدیو: {video}")
        files = _s(_cell(row, cols, "files")) or _s(_cell(row, cols, "note"))
        if files:
            note_parts.append("فایل‌های پیوست:\n" + "\n".join(f"• {x}" for x in files.split("\n")))
        parsed.lessons.append(
            LessonRow(
                where=where,
                course=course,
                ch_no=ch_no,
                ch_title=ch_title or f"فصل {ch_no}",
                code=_s(_cell(row, cols, "code")).upper(),
                title=title[:200],
                kind=kind,
                minutes=int(minutes) if minutes is not None else None,
                free=free,
                body=_s(_cell(row, cols, "body")),
                note="\n".join(note_parts),
                published=not status or status in PUBLISHED_STATUSES,
            )
        )


def _parse_bundles(ws, parsed: _Parsed, report: ImportReport):
    cols, rows = _read_table(ws, BUNDLE_COLS)
    if cols is None or "title" not in cols or "courses" not in cols:
        report.errors.append(f"برگه «{ws.title}»: ستون‌های «نام پکیج» و «دوره‌ها» لازم است.")
        return
    for number, row in rows:
        where = f"برگه «{ws.title}» ردیف {number}"
        title = _s(_cell(row, cols, "title"))
        codes = _split_codes(_cell(row, cols, "courses"))
        if not title or not codes:
            continue
        discount = _num(_cell(row, cols, "discount"))
        if discount is not None and discount > 1:
            discount /= 100
        access_cell = _cell(row, cols, "access")
        if _key(access_cell) in PERMANENT_WORDS and _s(access_cell) not in {"", "-", "—"}:
            access = LIFETIME_DAYS
        elif (n := _num(access_cell)) and n > 0:
            access = int(n)
        else:
            access = None
        order = _num(_cell(row, cols, "order"))
        price = _num(_cell(row, cols, "price"))
        parsed.bundles.append(
            BundleRow(
                where=where,
                code=_s(_cell(row, cols, "code")).upper(),
                title=title,
                courses=codes,
                price=int(price) if price is not None else -1,
                access_days=access,
                published_status=(
                    not _s(_cell(row, cols, "status"))
                    or _key(_cell(row, cols, "status")) in PUBLISHED_STATUSES
                ),
                summary=_s(_cell(row, cols, "summary"))[:300],
                description=_s(_cell(row, cols, "description")),
                cover=_s(_cell(row, cols, "cover")),
                slug=_s(_cell(row, cols, "slug")),
                order=int(order) if order is not None else len(parsed.bundles) + 1,
                discount=discount or 0,
            )
        )


def parse_workbook(fileobj, report: ImportReport) -> _Parsed:
    from openpyxl import load_workbook

    parsed = _Parsed()
    try:
        workbook = load_workbook(fileobj, data_only=True)
    except Exception:
        report.errors.append("فایل اکسل خوانده نشد؛ مطمئن شوید فرمت xlsx است.")
        return parsed
    sheets = [(ws, _sheet_role(ws.title)) for ws in workbook.worksheets]
    for ws, role in sheets:  # settings first: course prices may depend on them
        if role == "settings":
            _parse_settings(ws, parsed)
    for ws, role in sheets:
        if role == "categories":
            _parse_categories(ws, parsed, report)
    for ws, role in sheets:
        if role == "courses":
            _parse_courses(ws, parsed, report)
    for ws, role in sheets:
        if role == "lessons":
            _parse_lessons(ws, parsed, report)
    for ws, role in sheets:
        if role == "bundles":
            _parse_bundles(ws, parsed, report)
    if not (parsed.courses or parsed.lessons or parsed.bundles) and report.ok:
        report.errors.append("هیچ دوره، جلسه یا پکیجی در فایل پیدا نشد (نام برگه‌ها را بررسی کنید).")
    return parsed


# ── validation ───────────────────────────────────────────────────────────────
def _validate(parsed: _Parsed, report: ImportReport):
    known = {c.code: c for c in parsed.courses if c.code}
    db_codes = set(Course.objects.exclude(code="").values_list("code", flat=True))
    seen = set()
    for course in parsed.courses:
        if course.code and course.code in seen:
            report.errors.append(f"{course.where}: کد دوره {course.code} تکراری است.")
        seen.add(course.code)
        for code in course.prereqs:
            if code not in known and code not in db_codes:
                report.warnings.append(
                    f"{course.where}: پیش‌نیاز {code} پیدا نشد و نادیده گرفته شد."
                )
        if course.code in course.prereqs:
            report.errors.append(f"{course.where}: دوره نمی‌تواند پیش‌نیاز خودش باشد.")
    for lesson in parsed.lessons:
        if lesson.course not in known and lesson.course not in db_codes:
            report.errors.append(
                f"{lesson.where}: دوره {lesson.course} نه در فایل هست و نه در سیستم."
            )
    codes_seen = set()
    for lesson in parsed.lessons:
        if lesson.code:
            key = (lesson.course, lesson.code)
            if key in codes_seen:
                report.errors.append(f"{lesson.where}: کد جلسه {lesson.code} در دوره تکراری است.")
            codes_seen.add(key)
    prices = {c.code: c.price for c in parsed.courses if c.code}
    prices.update(
        {
            c: p
            for c, p in Course.objects.exclude(code="").values_list("code", "price")
            if c not in prices
        }
    )
    for bundle in parsed.bundles:
        missing = [c for c in bundle.courses if c not in prices]
        if missing:
            report.errors.append(f"{bundle.where}: دوره(های) {'، '.join(missing)} پیدا نشد.")
            continue
        if bundle.price < 0:
            total = sum(prices[c] for c in bundle.courses)
            bundle.price = _round(total * (1 - bundle.discount), 10_000)


# ── apply ────────────────────────────────────────────────────────────────────
def _free_flags(course: CourseRow, lessons: list[LessonRow]):
    rule = course.free_rule
    for index, lesson in enumerate(lessons):
        if lesson.free is not None:
            continue
        if rule[0] == "all":
            lesson.free = True
        elif rule[0] == "chapter":
            lesson.free = lesson.ch_no == rule[1]
        elif rule[0] == "first":
            lesson.free = index < rule[1]
        else:
            lesson.free = False


def _valid_image(data: bytes) -> bool:
    from PIL import Image

    try:
        Image.open(io.BytesIO(data)).verify()
    except Exception:
        return False
    return True


def _apply(parsed: _Parsed, report: ImportReport, covers: CoverSource | None, replace_covers: bool):
    from apps.shop.models import Bundle

    pending_covers: list[tuple] = []

    # categories (only the ones courses use)
    category_objs: dict[str, Category] = {}
    for course in parsed.courses:
        key = _key(course.category)
        if key in category_objs:
            continue
        row = parsed.categories.get(key)
        title = row.title if row else course.category
        obj = Category.objects.filter(title=title).first()
        created = obj is None
        if created:
            obj = Category(title=title)
            if row is None:
                report.warnings.append(
                    f"{course.where}: دسته «{title}» در برگه دسته‌ها نبود؛ ساخته شد."
                )
        if row:
            obj.description = row.description or obj.description
            obj.order = row.order
        obj.save()
        (report.created if created else report.updated)["دسته"] += 1
        category_objs[key] = obj
        category_objs[_key(title)] = obj

    # courses
    course_objs: dict[str, Course] = {}
    for row in parsed.courses:
        obj = Course.objects.filter(code=row.code).first() if row.code else None
        if obj is None:
            obj = Course.objects.filter(title=row.title).first()
        created = obj is None
        if created:
            obj = Course()
            base = row.slug or row.code.lower() or row.title
            obj.slug = _unique_slug(obj, base)
        obj.code = row.code
        obj.title = row.title
        obj.category = category_objs[_key(row.category)]
        obj.level = row.level
        obj.price = row.price
        obj.access_days = row.access_days
        obj.is_published = row.published
        obj.order = row.order
        if row.summary or created:
            obj.summary = row.summary
        if row.description or created:
            obj.description = row.description
        obj.save()
        (report.created if created else report.updated)["دوره"] += 1
        if row.code:
            course_objs[row.code] = obj
        if covers is not None and (created or not obj.cover or replace_covers):
            found = covers.find(row.cover, row.code)
            if found:
                pending_covers.append((obj, *found))
    for obj in Course.objects.exclude(code=""):
        course_objs.setdefault(obj.code, obj)

    for row in parsed.courses:
        if row.prereqs:
            obj = course_objs.get(row.code) or Course.objects.get(title=row.title)
            obj.prerequisites.set([course_objs[c] for c in row.prereqs if c in course_objs])

    # chapters + lessons
    by_course: dict[str, list[LessonRow]] = {}
    for lesson in parsed.lessons:
        by_course.setdefault(lesson.course, []).append(lesson)
    course_rows = {c.code: c for c in parsed.courses if c.code}
    for code, lessons in by_course.items():
        course = course_objs[code]
        if code in course_rows:
            _free_flags(course_rows[code], lessons)
        chapters: dict[int, Chapter] = {}
        counters: Counter = Counter()
        touched = set()
        for row in lessons:
            chapter = chapters.get(row.ch_no)
            if chapter is None:
                chapter = Chapter.objects.filter(course=course, order=row.ch_no).first()
                created = chapter is None
                if created:
                    chapter = Chapter(course=course, order=row.ch_no)
                chapter.title = row.ch_title
                chapter.save()
                (report.created if created else report.updated)["فصل"] += 1
                chapters[row.ch_no] = chapter
            lesson = None
            if row.code:
                lesson = Lesson.objects.filter(chapter__course=course, code=row.code).first()
            if lesson is None:
                lesson = Lesson.objects.filter(chapter=chapter, title=row.title, code="").first()
            created = lesson is None
            if created:
                lesson = Lesson()
            counters[row.ch_no] += 1
            lesson.chapter = chapter
            lesson.code = row.code
            lesson.title = row.title
            lesson.kind = row.kind
            lesson.duration_minutes = row.minutes
            lesson.is_free = bool(row.free)
            lesson.is_published = row.published
            lesson.order = counters[row.ch_no]
            lesson.internal_note = row.note
            if row.body or created:
                lesson.body = row.body
            lesson.save()
            touched.add(lesson.pk)
            (report.created if created else report.updated)["جلسه"] += 1
        extra = Lesson.objects.filter(chapter__course=course).exclude(pk__in=touched).count()
        if extra:
            report.warnings.append(
                f"دوره {code}: {extra} جلسه در سیستم هست که در اکسل نیست (دست‌نخورده ماند)."
            )

    # bundles
    for row in parsed.bundles:
        obj = Bundle.objects.filter(title=row.title).first()
        created = obj is None
        if created:
            obj = Bundle()
            obj.slug = _unique_slug(obj, row.slug or row.code.lower() or row.title)
        obj.title = row.title
        obj.price = row.price
        obj.access_days = row.access_days
        obj.order = row.order
        members = [course_objs[c] for c in row.courses]
        obj.is_published = row.published_status and all(m.is_published for m in members)
        if row.summary or created:
            obj.summary = row.summary
        if row.description or created:
            obj.description = row.description
        obj.save()
        obj.courses.set(members)
        (report.created if created else report.updated)["پکیج"] += 1
        if covers is not None and (created or not obj.cover or replace_covers):
            found = covers.find(row.cover, row.code)
            if found:
                pending_covers.append((obj, *found))
    return pending_covers


def import_workbook(
    fileobj,
    *,
    covers: CoverSource | None = None,
    dry_run: bool = False,
    replace_covers: bool = False,
) -> ImportReport:
    report = ImportReport(dry_run=dry_run)
    parsed = parse_workbook(fileobj, report)
    if report.ok:
        _validate(parsed, report)
    if not report.ok:
        return report

    try:
        with transaction.atomic():
            pending = _apply(parsed, report, covers, replace_covers)
            if dry_run:
                transaction.set_rollback(True)
    except Exception as exc:  # defensive: report instead of a 500 in admin
        report.errors.append(f"خطای غیرمنتظره هنگام ذخیره: {exc}")
        return report

    for obj, name, data in pending:
        if not _valid_image(data):
            report.warnings.append(f"تصویر «{name}» معتبر نیست و نادیده گرفته شد.")
            continue
        if not dry_run:
            obj.cover.save(name, ContentFile(data), save=True)
        report.created["تصویر"] += 1
    return report
