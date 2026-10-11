"""Builds the downloadable sample workbook for the catalog importer."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
INPUT_FILL = PatternFill("solid", fgColor="FEF9C3")


def _sheet(wb, title, headers, rows, widths, lists=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = "A2"
    for col, options in (lists or {}).items():
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"')
        validation.add(f"{col}2:{col}500")
        ws.add_data_validation(validation)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def build_template() -> bytes:
    wb = Workbook()
    guide = wb.active
    guide.title = "راهنما"
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 110
    guide.append(["بخش", "توضیح"])
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    notes = [
        (
            "کلیات",
            "برگه‌های «دسته‌ها»، «دوره‌ها»، «درس‌ها» و «پکیج‌ها» را پر کنید. نام برگه‌ها و "
            "ستون‌ها نباید تغییر کند (ستون‌های اختیاری را می‌توان خالی گذاشت).",
        ),
        (
            "دوره‌ها",
            "ستون «کد» یکتاست (مثل F1 یا P2) و همان کلید به‌روزرسانی است: ایمپورت دوباره "
            "همان کد را ویرایش می‌کند. «دسته» کد دسته (C1) یا نام آن است.",
        ),
        (
            "قیمت",
            "در «قیمت (تومان)» عدد بنویسید؛ ۰ یعنی رایگان. اگر برگه «تنظیمات قیمت» "
            "(قیمت پایه و ضریب رده‌ها) داشته باشید می‌توانید به‌جای قیمت، «رده قیمت» بنویسید.",
        ),
        ("دسترسی (روز)", "عدد روز، یا «نامحدود» برای دسترسی دائمی."),
        ("پیش‌نیاز", "کد دوره‌ها با ویرگول، مثلاً F2، F3 (فقط برای نمایش به دانشجو)."),
        (
            "درس رایگان",
            "همه | هیچ | فصل ۱ | ۲ درس (دو درس اول). اگر در برگه درس‌ها ستون «رایگان» "
            "برای یک جلسه پر باشد (بله/خیر) همان ملاک است.",
        ),
        (
            "وضعیت",
            "دوره/جلسه فقط وقتی «منتشر» می‌شود که وضعیتش یکی از اینها باشد: منتشر، "
            "محتوا آماده، آماده، آپلود شد. غیر از این پیش‌نویس می‌ماند.",
        ),
        (
            "درس‌ها",
            "هر ردیف یک جلسه: کد دوره، شماره و عنوان فصل، کد جلسه (L01…)، عنوان، نوع، "
            "مدت (دقیقه). «نام فایل ویدیو» و «فایل‌های پیوست» فقط یادداشت داخلی‌اند "
            "(فایل واقعی را بعداً از صفحه جلسه آپلود کنید).",
        ),
        (
            "برگه درس‌ها",
            "می‌توانید برای هر دوره یک برگه جدا با نام «درس‌های F2» بسازید؛ آن‌وقت "
            "ستون «دوره» لازم نیست.",
        ),
        (
            "پکیج‌ها",
            "«دوره‌ها» کد دوره‌هاست. قیمت پکیج را بنویسید، یا «تخفیف» (مثلاً ۰٫۳ یا ۳۰%) "
            "بدهید تا از جمع قیمت دوره‌ها حساب شود. پکیج فقط وقتی منتشر می‌شود که همه "
            "دوره‌هایش منتشر باشند. دسترسی: «مادام‌العمر» یا تعداد روز.",
        ),
        (
            "تصویر",
            "ستون «تصویر» نام فایل است و تصاویر را به‌صورت zip در همان صفحه ایمپورت بدهید؛ "
            "یا فایل‌ها را با نام کد (F2.png) بنامید تا خودکار پیدا شوند.",
        ),
    ]
    for row in notes:
        guide.append(list(row))
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    _sheet(
        wb,
        "دسته‌ها",
        ["کد دسته", "نام دسته", "توضیح"],
        [["C1", "مبانی و پیش‌نیازها", "دوره‌های شروع مسیر"]],
        [12, 36, 50],
    )
    _sheet(
        wb,
        "دوره‌ها",
        [
            "کد",
            "عنوان دوره",
            "دسته",
            "سطح",
            "پیش‌نیاز",
            "قیمت (تومان)",
            "دسترسی (روز)",
            "درس رایگان",
            "خلاصه",
            "توضیحات",
            "تصویر",
            "وضعیت",
        ],
        [
            [
                "S1",
                "نمونه: مقدمه‌ای بر GeoAI",
                "C1",
                "مقدماتی",
                "",
                0,
                "نامحدود",
                "همه",
                "یک خلاصه کوتاه (حداکثر ۳۰۰ حرف)",
                "توضیح کامل دوره…",
                "S1.png",
                "منتشر",
            ],
            [
                "S2",
                "نمونه: دوره پولی",
                "C1",
                "متوسط",
                "S1",
                1500000,
                365,
                "فصل ۱",
                "خلاصه",
                "توضیح کامل…",
                "",
                "برنامه‌ریزی",
            ],
        ],
        [8, 36, 10, 12, 14, 16, 14, 14, 40, 50, 14, 14],
        lists={"D": ["مقدماتی", "متوسط", "پیشرفته"]},
    )
    _sheet(
        wb,
        "درس‌ها",
        [
            "دوره",
            "شماره فصل",
            "فصل",
            "کد درس",
            "عنوان درس",
            "نوع",
            "مدت (دقیقه)",
            "رایگان",
            "توضیح",
            "نام فایل ویدیو",
            "فایل‌های پیوست",
            "وضعیت",
        ],
        [
            [
                "S1",
                1,
                "شروع",
                "L01",
                "معرفی دوره",
                "ویدیو",
                8,
                "بله",
                "در این جلسه…",
                "S1-CH1-L01.mp4",
                "scripts/S1-CH1-L01.md",
                "محتوا آماده",
            ],
            [
                "S1",
                1,
                "شروع",
                "L02",
                "نصب ابزارها",
                "ویدیو",
                12,
                "",
                "",
                "S1-CH1-L02.mp4",
                "",
                "محتوا آماده",
            ],
        ],
        [8, 12, 22, 10, 36, 12, 12, 10, 40, 22, 36, 14],
        lists={
            "F": ["ویدیو", "ارائه HTML", "PDF / اسلاید", "نوت‌بوک", "داده مکانی", "کد", "متن"],
            "H": ["بله", "خیر"],
        },
    )
    _sheet(
        wb,
        "پکیج‌ها",
        [
            "کد",
            "نام پکیج",
            "دوره‌ها",
            "تخفیف",
            "قیمت پکیج (تومان)",
            "دسترسی",
            "خلاصه",
            "توضیحات",
            "تصویر",
            "وضعیت",
        ],
        [
            [
                "B1",
                "نمونه: مسیر شروع",
                "S1، S2",
                0.3,
                "",
                "مادام‌العمر",
                "خلاصه پکیج",
                "توضیح پکیج…",
                "B1.png",
                "منتشر",
            ]
        ],
        [8, 30, 18, 10, 20, 14, 36, 50, 14, 14],
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
